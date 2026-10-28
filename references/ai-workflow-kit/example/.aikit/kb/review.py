"""Review the whole KB store in one spreadsheet: export → the reviewer marks Keep/Remove/Change → import.

Export writes one row per store record with a Decision column (Keep by default) and a Note.
Import refuses a sheet whose group documents changed after the export, proposes a `remove` for
every Remove row, and lists Change rows for the AI to draft as edits. With `approve_as` (only after
the reviewer said OK in chat) it also reviews and applies the removals; the ledger keeps every removed row.
Needs openpyxl.
"""

from datetime import datetime, timezone
from pathlib import Path

DECISIONS = ('Keep', 'Remove', 'Change')
COLUMNS = ['Group', 'Scope', 'ID', 'Type', 'Status', 'Section', 'Text', 'Citation', 'Decision', 'Note']
WIDTHS = [18, 9, 22, 7, 12, 28, 90, 40, 11, 40]


def _openpyxl(kb):
    try:
        import openpyxl
    except ImportError as e:
        raise kb.KBError('The review sheet needs openpyxl (pip install openpyxl)') from e
    return openpyxl


def _store_groups(kb, root, names=None):
    groups = [g for g in kb.load_groups(root) if g['mode'] == 'store']
    if names:
        known = {g['group']: g for g in kb.load_groups(root)}
        for name in names:
            if name not in known:
                raise kb.KBError(f'Unknown group: {name}')
            if known[name]['mode'] != 'store':
                raise kb.KBError(f'{name} is read-only; its source document is reviewed where it lives')
        groups = [g for g in groups if g['group'] in names]
    if not groups:
        raise kb.KBError('No store group to review')
    return groups


def _document_sha(kb, root, cfg):
    import store
    _, _, state = store.load(kb, root, cfg)
    return store.digest(state['document'].encode('utf-8'))


def export(kb, root, out, names=None):
    openpyxl = _openpyxl(kb)
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    groups = _store_groups(kb, root, names)
    _, records, _ = kb.validated_records(root, strict=False)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Records'
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D6E2EF')
    count = 0
    for cfg in groups:
        rows = [r for r in records if r['group'] == cfg['group']]
        if len(rows) != len(cfg['expected_ids']) and not rows and cfg['expected_ids']:
            raise kb.KBError(f'{cfg["group"]} is unavailable; run lint')
        for r in sorted(rows, key=lambda r: r.get('source_line', 0)):
            cite = kb.citation(r)['citation']
            sheet.append([cfg['group'], cfg['scope'], r['id'], r['type'], r['status'],
                          r['section'].lstrip('#').strip(), r['text'].strip(), cite, 'Keep', ''])
            count += 1
    last = sheet.max_row
    choice = DataValidation(type='list', formula1='"' + ','.join(DECISIONS) + '"', allow_blank=False,
                            showErrorMessage=True, error='Choose Keep, Remove or Change')
    sheet.add_data_validation(choice)
    if last > 1:
        choice.add(f'I2:I{last}')
    for n, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(n)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = f'A1:J{last}'

    summary = book.create_sheet('Summary')
    summary.append(['Group', 'Records', 'Remove', 'Change'])
    for n, cfg in enumerate(groups, 2):
        g = f'A{n}'
        summary.append([cfg['group'], f'=COUNTIF(Records!A:A,{g})',
                        f'=COUNTIFS(Records!A:A,{g},Records!I:I,"Remove")',
                        f'=COUNTIFS(Records!A:A,{g},Records!I:I,"Change")'])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.column_dimensions['A'].width = 22

    meta = book.create_sheet('Meta')
    meta.append(['Group', 'Document SHA256', 'Source'])
    for cfg in groups:
        meta.append([cfg['group'], _document_sha(kb, root, cfg), cfg['source']])
    meta.append([])
    meta.append(['Exported (UTC)', datetime.now(timezone.utc).isoformat()])
    meta.append(['How to use', 'Set Decision to Remove or Change and say why in Note. Then '
                 '`kb.py review import --file <this file>`. Do not edit the other columns.'])
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    return {'file': str(out), 'records': count, 'groups': [g['group'] for g in groups]}


def _read(kb, path):
    openpyxl = _openpyxl(kb)
    try:
        book = openpyxl.load_workbook(path, data_only=True)
        sheet, meta = book['Records'], book['Meta']
    except (OSError, KeyError, ValueError) as e:
        raise kb.KBError(f'Not a KB review sheet: {path}') from e
    header = [c.value for c in sheet[1]]
    if header[:len(COLUMNS)] != COLUMNS:
        raise kb.KBError('The Records sheet columns were changed; export a new sheet')
    shas = {}
    for row in meta.iter_rows(min_row=2, values_only=True):
        if row and row[0] and row[1] and len(str(row[1])) == 64:
            shas[row[0]] = row[1]
    rows = []
    for n, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(row):
            continue
        item = dict(zip(COLUMNS, row))
        decision = str(item['Decision'] or '').strip().capitalize()
        if decision not in DECISIONS:
            raise kb.KBError(f'Decision must be Keep, Remove or Change (row {n}: {item["Decision"]!r})')
        rows.append({'row': n, 'group': item['Group'], 'id': item['ID'], 'decision': decision,
                     'note': str(item['Note'] or '').strip()})
    return rows, shas


def import_(kb, root, path, approve_as=None, reason=None):
    import store
    path = Path(path)
    rows, shas = _read(kb, path)
    sheet_sha = kb.sha(path.read_bytes())
    groups = {g['group']: g for g in _store_groups(kb, root)}
    acted = [r for r in rows if r['decision'] != 'Keep']
    # Check everything before the first proposal, so a bad sheet changes nothing.
    for r in acted:
        cfg = groups.get(r['group'])
        if cfg is None or r['group'] not in shas:
            raise kb.KBError(f'Row {r["row"]}: {r["group"]} is not a store group of this sheet')
        if _document_sha(kb, root, cfg) != shas[r['group']]:
            raise kb.KBError(f'{r["group"]} changed since the export; export a new sheet and mark it again')
        if r['id'] not in {x['id'] for x in kb.validated_group(root, cfg)[1]}:
            raise kb.KBError(f'Row {r["row"]}: {r["id"]} is not in {r["group"]}')
    if approve_as is not None:
        store.nonempty(approve_as, 'Approver')
        store.nonempty(reason, 'Reason')
    proposed, applied = [], []
    for r in [r for r in acted if r['decision'] == 'Remove']:
        status = store.status(kb, root, r['group'])
        pending = [p for p in status['proposals'] if p['op'] == 'remove' and p['id'] == r['id']
                   and p['review'] == 'pending']
        if pending:
            number = pending[0]['proposal']
        else:
            why = f'KB review sheet {path.name} (sha256 {sheet_sha[:12]}): ' + (r['note'] or 'marked Remove')
            number = store.propose(kb, root, 'remove', r['id'], None, '', why, 'AI', group=r['group'])['proposal']
        proposed.append({'group': r['group'], 'id': r['id'], 'proposal': number, 'note': r['note']})
    if approve_as is not None:
        for p in proposed:
            store.review(kb, root, p['proposal'], 'approved', approve_as,
                         f'{reason} (review sheet {path.name}, sha256 {sheet_sha[:12]})', group=p['group'])
            store.apply(kb, root, p['proposal'], 'AI', group=p['group'])
            applied.append(p)
    changes = [{'group': r['group'], 'id': r['id'], 'note': r['note']} for r in acted if r['decision'] == 'Change']
    return {'sheet': path.name, 'sheet_sha256': sheet_sha, 'proposed': proposed, 'applied': applied,
            'change_requests': changes,
            'next': None if approve_as or not proposed else
            f'Removals are proposed, not applied. After {kb.kbpaths.reviewer(root)} says OK in chat, run import again '
            f'with --approve-as "{kb.kbpaths.reviewer(root)}" --reason "<quote the OK>".'}
