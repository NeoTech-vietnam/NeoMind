"""Read-only parsers (md-table, md-bullet-id, docx-field-table, xlsx-table) on synthetic sources, and
cross-checks of the DOCX and XLSX readers against independent libraries when those are installed."""

import re
import unittest

from . import project
from .project import ProjectCase, entry, kb, para, table

try:
    import docx as python_docx  # Optional independent oracle for the DOCX parser.
except ImportError:  # pragma: no cover - depends on the environment
    python_docx = None
try:
    import openpyxl  # Optional independent oracle for the XLSX parser.
except ImportError:  # pragma: no cover - depends on the environment
    openpyxl = None

MD_TABLE = project.NOTES_MD
BULLETS = project.RULES_MD


def cfg_md_table(**extra):
    return entry('ops-exceptions', **extra)


def cfg_bullets(**extra):
    return entry('ops-rules', **extra)


def cfg_docx(**extra):
    return entry('ops-spec', **extra)


def cfg_xlsx(**extra):
    return entry('ops-backlog', **extra)


class ParserCase(ProjectCase):
    """A project with no groups: each test writes its own source and builds the group entry it needs."""

    GROUPS = ()
    SETUP = False

    def registry(self, groups):
        self.p.write_registry(groups)


class MdTableTests(ParserCase):
    def test_rows_of_the_named_section_only_verbatim(self):
        self.write('docs/notes.md', MD_TABLE)
        records = kb.extract(self.root, cfg_md_table())
        self.assertEqual([r['id'] for r in records], ['OPS-EXC-1', 'OPS-EXC-2'])
        lines = MD_TABLE.splitlines(keepends=True)
        for r in records:
            self.assertEqual(r['text'], lines[r['source_line'] - 1])
            self.assertEqual((r['type'], r['status'], r['approval_status']),
                             ('CODE', 'as-built', 'not-stated-in-source'))
            self.assertEqual(r['section'], '## 3. Delivery Exceptions')
        self.assertEqual(records[0]['source_id'], '1')
        self.assertEqual(records[1]['links'], [])  # Bare numbers are not links.

    def test_unexpected_row_duplicate_and_missing_section_fail(self):
        cfg = cfg_md_table()
        for text, message in [(MD_TABLE.replace('| 2 | damaged', '| x | damaged'), 'Unexpected row'),
                              (MD_TABLE.replace('| 2 | damaged', '| 1 | damaged'), 'Duplicate source ID'),
                              (MD_TABLE.replace('## 3. Delivery', '## 4. Delivery'), 'Section not found')]:
            with self.subTest(message=message):
                self.write('docs/notes.md', text)
                with self.assertRaisesRegex(kb.KBError, message):
                    kb.extract(self.root, cfg)

    def test_deferred_id_is_skipped_and_reported(self):
        self.write('docs/notes.md', MD_TABLE.replace('| 2 | damaged, see 1 |', '| 2 | damaged, see 1 |\n| 1 | again |'))
        cfg = cfg_md_table(deferred_ids={'1': 'Duplicate in source; owner to fix'}, expected_ids=['OPS-EXC-2'])
        records, notes = kb.extract_with_notes(self.root, cfg)
        self.assertEqual([r['id'] for r in records], ['OPS-EXC-2'])
        self.assertEqual(notes['deferred_ids'], {'1': 'Duplicate in source; owner to fix'})
        with self.assertRaises(kb.KBError):  # An empty group is not a group.
            kb.checked_group(self.root, {**cfg, 'expected_ids': []})

    def test_bold_ids_and_links_for_named_ids(self):
        text = '## Decisions\n\n| # | Decision |\n| --- | --- |\n| **K1** | Inputs |\n| K2 | See K1 and D4 |\n'
        self.write('docs/kb.md', text)
        cfg = cfg_md_table(source='docs/kb.md', section='## Decisions', id_pattern=r'K\d+', id_prefix='',
                           type='DEC', status='decided', approval_status='source-reported-decided',
                           expected_ids=['K1', 'K2'])
        records = {r['id']: r for r in kb.extract(self.root, cfg)}
        self.assertEqual(records['K1']['source_id'], 'K1')
        self.assertEqual(records['K2']['links'], ['K1'])

    def test_config_validation(self):
        good = cfg_md_table()
        bad_cases = [('id_pattern', '('), ('type', 'XYZ'), ('status', 'maybe'),
                     ('approval_status', 'approved'), ('section', ''), ('id_prefix', 'bad prefix '),
                     ('expected_ids', ['FAULT-1']), ('mode', 'store'), ('deferred_ids', {'1': ''})]
        for key, value in bad_cases:
            with self.subTest(key=key):
                cfg = {**good, key: value}
                if key == 'mode':
                    cfg.update(authority='kb-store', store='knowledge/kb/store/x.jsonl')
                with self.assertRaises(kb.KBError):
                    kb.checked_group(self.root, cfg)
        kb.checked_group(self.root, good)

    def test_rows_match_an_independent_reading_of_the_code_table(self):
        # The exception codes live in code too; the document and the code must name the same set.
        enum = ('class ExceptionCode:\n    NONE = 0\n    ADDRESS_NOT_FOUND = 1\n    DAMAGED = 2\n')
        self.write('docs/notes.md', MD_TABLE)
        codes = sorted(int(v) for v in re.findall(r'=\s*(\d+)', enum) if int(v))
        cfg = cfg_md_table(expected_ids=[f'OPS-EXC-{c}' for c in codes])
        records = kb.extract(self.root, cfg)
        self.assertEqual([r['source_id'] for r in records], [str(c) for c in codes])
        lines = MD_TABLE.splitlines(keepends=True)
        for r in records:
            self.assertEqual(r['text'], lines[r['source_line'] - 1])


class BulletTests(ParserCase):
    def test_rule_blocks_stop_at_blank_heading_or_next_bullet(self):
        self.write('docs/rules.md', BULLETS)
        records = {r['id']: r for r in kb.extract(self.root, cfg_bullets())}
        lines = BULLETS.splitlines(keepends=True)
        self.assertEqual(records['R0']['text'], ''.join(lines[2:4]))
        self.assertEqual(records['R11']['text'], ''.join(lines[7:10]))
        self.assertEqual(records['R1']['text'], lines[10])
        self.assertEqual(records['R2']['text'], lines[12])
        self.assertEqual((records['R11']['source_line'], records['R11']['source_end_line']), (8, 10))
        self.assertEqual(records['R11']['links'], ['R1'])
        self.assertEqual(records['R11']['section'], '## Meta')
        self.assertEqual(records['R0']['section'], '# RULES')
        self.assertEqual(records['R2']['status'], 'in-force')

    def test_duplicate_and_unexpected_ids_fail(self):
        self.write('docs/rules.md', BULLETS + '- **R1** — Again.\n')
        with self.assertRaisesRegex(kb.KBError, 'Duplicate source ID: R1'):
            kb.extract(self.root, cfg_bullets())
        self.write('docs/rules.md', BULLETS + '- **R30** — New.\n')
        with self.assertRaisesRegex(kb.KBError, 'unexpected=.*R30'):
            kb.extract(self.root, cfg_bullets())

    def test_decisions_and_rules_are_verbatim(self):
        # Verbatim check against an independent line scan of the same text, for a table and for bullets.
        decided = project.DECISIONS_MD
        self.write('docs/decisions.md', decided)
        cases = [('docs/decisions.md',
                  cfg_md_table(source='docs/decisions.md', section='## Decided', id_pattern=r'D\d+', id_prefix='',
                               type='DEC', status='decided', approval_status='source-reported-decided'),
                  r'^\|\s*(D\d+)\s*\|', 4),
                 ('docs/rules.md', cfg_bullets(), r'^(?:- |> )\*\*(R\d+)\b', 4)]
        self.write('docs/rules.md', BULLETS)
        for rel, cfg, pattern, count in cases:
            with self.subTest(source=rel):
                lines = self.p.path(rel).read_bytes().replace(b'\r\n', b'\n').decode('utf-8-sig') \
                    .splitlines(keepends=True)
                ids = [m[1] for m in (re.match(pattern, line) for line in lines) if m]
                if rel == 'docs/decisions.md':
                    ids = ids[:count]  # D1-D4 sit under `## Decided`; D5 is in another table.
                cfg['expected_ids'] = ids
                records = kb.extract(self.root, cfg)
                self.assertEqual(sorted(r['id'] for r in records), sorted(ids))
                for r in records:
                    end = r.get('source_end_line', r['source_line'])
                    self.assertEqual(r['text'], ''.join(lines[r['source_line'] - 1:end]))


class DocxTests(ParserCase):
    def test_field_tables_become_records_with_heading_path(self):
        self.write('docs/spec.docx', project.docx_bytes(project.spec_body()))
        records, notes = kb.extract_with_notes(self.root, cfg_docx())
        r = {x['id']: x for x in records}
        fr = r['SPEC-FUN-SCAN-001']
        self.assertEqual(fr['text'], 'ID: SPEC-FUN-SCAN-001\nTitle: Scan feedback\nPriority: High\n'
                                     'Dependencies: - Preceding: SPEC-FUN-SCAN-002\n- Logging\n'
                                     'Status: Active')
        self.assertEqual((fr['status'], fr['title'], fr['source_locator']), ('active', 'Scan feedback', 'table 2'))
        self.assertEqual(fr['section'], 'Detailed Functional Requirements > Stage01-Scanning > Scan feedback')
        self.assertEqual(fr['links'], ['SPEC-FUN-SCAN-002'])
        self.assertNotIn('source_line', fr)
        nfr = r['SPEC-NFR-CMN-001']
        self.assertEqual((nfr['status'], nfr['section']), ('unknown', 'Non-Functional Requirements > Performance'))
        self.assertEqual(notes['skipped_source_ids'], [{'table': 1, 'id': project.TEMPLATE[1][1]}])
        self.assertEqual(kb.citation(fr)['citation'], 'docs/spec.docx#table-2')

    def test_tracked_changes_accept_insertions_and_drop_deletions(self):
        row = ('<w:tr><w:tc>' + para('Purpose') + '</w:tc><w:tc><w:p><w:r><w:t>Play </w:t></w:r>'
               '<w:del w:id="1"><w:r><w:tab/><w:delText>Single</w:delText></w:r></w:del>'
               '<w:ins w:id="2"><w:r><w:t>Three</w:t></w:r></w:ins><w:r><w:tab/><w:t>beeps</w:t>'
               '<w:br/><w:t>next</w:t><w:noBreakHyphen/><w:t>line</w:t></w:r></w:p></w:tc></w:tr>')
        tracked = table(project.NFR_1[:3]).replace('</w:tbl>', row + '</w:tbl>')
        self.write('docs/spec.docx',
                   project.docx_bytes(project.spec_body().replace(table(project.NFR_1), tracked)))
        records = {r['id']: r for r in kb.extract(self.root, cfg_docx())}
        record = records['SPEC-NFR-CMN-001']
        self.assertTrue(record['text'].endswith('Purpose: Play Three\tbeeps\nnext-line'), record['text'])
        self.assertTrue(record['tracked_changes'])
        self.assertNotIn('tracked_changes', records['SPEC-FUN-SCAN-001'])

    def test_duplicate_template_corrupt_and_deferred(self):
        self.write('docs/spec.docx', project.docx_bytes(project.spec_body(table(project.FUN_1))))
        with self.assertRaisesRegex(kb.KBError, 'Duplicate source ID: SPEC-FUN-SCAN-001'):
            kb.extract(self.root, cfg_docx())
        cfg = cfg_docx(deferred_ids={'SPEC-FUN-SCAN-001': 'Same ID in tables 2 and 5'},
                       expected_ids=['SPEC-NFR-CMN-001'])
        records, notes = kb.extract_with_notes(self.root, cfg)
        self.assertEqual([r['id'] for r in records], ['SPEC-NFR-CMN-001'])
        self.assertIn('SPEC-FUN-SCAN-001', notes['deferred_ids'])
        self.write('docs/spec.docx', b'not a zip')
        with self.assertRaisesRegex(kb.KBError, 'not a readable DOCX'):
            kb.extract(self.root, cfg_docx())

    def test_changed_docx_is_stale_and_reimport_clears(self):
        self.write('docs/spec.docx', project.docx_bytes(project.spec_body()))
        self.registry([cfg_docx()])
        kb.import_source(self.root, 'ops-spec')
        self.assertEqual(kb.get(self.root, 'SPEC-NFR-CMN-001')['scope'], 'ops')
        self.write('docs/spec.docx', project.docx_bytes(project.spec_body().replace('200 ms', '100 ms')))
        with self.assertRaisesRegex(kb.KBError, 'STALE'):
            kb.get(self.root, 'SPEC-NFR-CMN-001')
        kb.import_source(self.root, 'ops-spec')
        self.assertIn('100 ms', kb.get(self.root, 'SPEC-NFR-CMN-001')['text'])

    @unittest.skipUnless(python_docx, 'python-docx not installed')
    def test_spec_matches_an_independent_docx_reader(self):
        document = python_docx.Document()
        document.add_heading('Detailed Functional Requirements', 1)
        for rows in (project.TEMPLATE, project.FUN_1, project.NFR_1, project.FUN_1):
            grid = document.add_table(rows=len(rows), cols=2)
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    grid.cell(r, c).text = value
        path = self.p.path('docs/spec.docx')
        document.save(str(path))
        expected = {}
        for tbl in python_docx.Document(str(path)).tables:
            rows = [[c.text.strip() for c in row.cells] for row in tbl.rows]
            fields = [row for row in rows[1:] if any(row)]  # Some tables carry an empty spacer row.
            ids = [row[1] for row in fields if row[0] == 'ID']
            if rows and rows[0][0] == 'Field' and ids:
                expected.setdefault(ids[0], []).append(fields)
        dup = sorted(k for k, v in expected.items() if len(v) > 1)
        pattern = cfg_docx()['id_pattern']
        good = [k for k in expected if re.fullmatch(pattern, k) and k not in dup]
        self.assertEqual((dup, good), (['SPEC-FUN-SCAN-001'], ['SPEC-NFR-CMN-001']))
        cfg = cfg_docx(deferred_ids={k: 'Duplicate in source' for k in dup}, expected_ids=good)
        records, notes = kb.extract_with_notes(self.root, cfg)
        self.assertEqual(len(records), len(good))
        self.assertEqual(len(notes['skipped_source_ids']),
                         len([k for k in expected if not re.fullmatch(pattern, k)]))
        for r in records:
            rows = expected[r['source_id']][0]
            self.assertEqual(r['text'], '\n'.join(f'{a}: {b}' for a, b, *_ in rows), r['id'])


class XlsxTests(ParserCase):
    def test_rows_under_the_id_header_become_records(self):
        self.write('docs/backlog.xlsx', project.xlsx_bytes({'Summary': [['Sprint']], 'Backlog': project.BACKLOG_ROWS}))
        records, notes = kb.extract_with_notes(self.root, cfg_xlsx())
        r = {x['id']: x for x in records}
        self.assertEqual(r['B-01']['text'], 'ID: B-01\nItem: Scanner set-up\nStatus: Closed\nRemaining: -')
        self.assertEqual(r['B-02']['text'], 'ID: B-02\nItem: Late-parcel accuracy\nStatus: Partly closed\n'
                                            'Remaining: About 4% gap.\nSecond line\nPriority: 2')
        self.assertEqual((r['B-01']['status'], r['B-02']['status']), ('closed', 'partly-closed'))
        self.assertEqual((r['B-02']['title'], r['B-02']['source_locator']),
                         ('Late-parcel accuracy', 'sheet Backlog row 6'))
        # A row without an ID names the part of the table below it.
        self.assertEqual((r['B-01']['section'], r['B-02']['section']),
                         ('Backlog', 'Backlog > Section note without an ID'))
        self.assertEqual(kb.citation(r['B-02'])['citation'], 'docs/backlog.xlsx#sheet-Backlog-row-6')
        self.assertEqual(notes['skipped_source_ids'], [{'row': 7, 'id': 'Total'}])

    def test_done_and_partly_done_statuses(self):
        rows = [['Item #', 'Description', 'Status'], ['', 'Device work'], ['1', 'Probing', 'Done'],
                ['2', 'Scheduling', 'Partly done'], ['3', 'Alarms', 'Not planned']]
        self.write('docs/backlog.xlsx', project.xlsx_bytes({'Items': rows}))
        cfg = cfg_xlsx(sheet='Items', id_column='Item #', id_pattern=r'\d+', id_prefix='OPS-BKL-',
                       title_field='Description', expected_ids=['OPS-BKL-1', 'OPS-BKL-2', 'OPS-BKL-3'])
        status = {r['id']: (r['status'], r['section']) for r in kb.extract(self.root, cfg)}
        self.assertEqual(status, {'OPS-BKL-1': ('done', 'Items > Device work'),
                                  'OPS-BKL-2': ('partly-done', 'Items > Device work'),
                                  'OPS-BKL-3': ('unknown', 'Items > Device work')})

    def test_missing_sheet_header_duplicate_and_corrupt(self):
        rows = project.BACKLOG_ROWS
        cases = [({'Other': rows}, 'Sheet not found: Backlog'),
                 ({'Backlog': [r for r in rows if r[:1] != ['ID']]}, 'No header row'),
                 ({'Backlog': rows + [['B-01', 'Again']]}, 'Duplicate source ID: B-01')]
        repeated = [list(r) for r in rows]
        repeated[2][4] = 'Item'
        cases.append(({'Backlog': repeated}, 'Repeated column header'))
        for sheets, message in cases:
            with self.subTest(message=message):
                self.write('docs/backlog.xlsx', project.xlsx_bytes(sheets))
                with self.assertRaisesRegex(kb.KBError, message):
                    kb.extract(self.root, cfg_xlsx())
        self.write('docs/backlog.xlsx', b'PK not really')
        with self.assertRaisesRegex(kb.KBError, 'not a readable XLSX'):
            kb.extract(self.root, cfg_xlsx())
        with self.assertRaises(kb.KBError):
            kb.checked_group(self.root, cfg_xlsx(sheet=''))

    @unittest.skipUnless(openpyxl, 'openpyxl not installed')
    def test_backlog_matches_an_independent_xlsx_reader(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Backlog'
        for row in project.BACKLOG_ROWS:
            sheet.append(row)
        book.save(self.p.path('docs/backlog.xlsx'))
        sheet = openpyxl.load_workbook(self.p.path('docs/backlog.xlsx'), data_only=True)['Backlog']
        rows = [['' if v is None else str(v).strip() for v in row] for row in sheet.iter_rows(values_only=True)]
        start = next(n for n, row in enumerate(rows) if 'ID' in row)
        header = rows[start]
        stop = header.index('', header.index('ID')) if '' in header[header.index('ID'):] else len(header)
        header = header[:stop]  # The backlog table ends at the first empty header cell.
        oracle = {}
        for row in rows[start + 1:]:
            if re.fullmatch(r'B-\d+', row[header.index('ID')]):
                oracle[row[header.index('ID')]] = '\n'.join(f'{h}: {v}' for h, v in zip(header, row) if h and v)
        self.assertNotIn('Count', header)
        self.assertEqual(sorted(oracle), ['B-01', 'B-02'])
        records = kb.extract(self.root, cfg_xlsx(expected_ids=list(oracle)))
        self.assertEqual({r['id']: r['text'] for r in records}, oracle)


if __name__ == '__main__':
    unittest.main()
